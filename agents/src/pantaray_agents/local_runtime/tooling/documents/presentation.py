"""Render a .pptx deck as one section per slide, using python-pptx.

A slide is partly a picture of itself, so the text below cannot be all of what
it carries: its words, tables and chart data are written out, while where its
shapes sit stays behind in the picture, and the notes say which is which.
"""

from __future__ import annotations

from bisect import bisect_right
from collections import Counter
from collections.abc import Iterator, Sequence
from itertools import accumulate, zip_longest
from typing import BinaryIO, Final

from pptx import Presentation
from pptx.chart.chart import Chart
from pptx.chart.plot import PlotFactory
from pptx.oxml.chart.plot import BaseChartElement
from pptx.oxml.chart.series import CT_SeriesComposite
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import BaseOxmlElement
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture
from pptx.slide import Slide
from pptx.table import Table, _Row

from .document_model import (
    MAX_TABLE_COLUMNS,
    MAX_TABLE_ROWS,
    DocumentChart,
    DocumentImage,
    DocumentTextBudget,
    ExtractedDocument,
    pipe_table,
    resolve_start_unit,
    stored_value_text,
    table_cut_notes,
)

# python-pptx renders a soft line break inside a paragraph as a vertical tab.
_LINE_BREAK: Final = "\v"
# ECMA-376 types a paragraph's lvl as 0 through 8, and python-pptx returns the
# attribute unchecked when it reads a file.
_MAX_OUTLINE_LEVEL: Final = 8
# PowerPoint writes show="0" for a slide it skips in a show; ECMA-376 types the
# attribute as xsd:boolean, so another producer may write the word instead.
_HIDDEN_SLIDE_VALUES: Final = frozenset({"0", "false"})
# What python-pptx raises for a chart it declines to describe, measured by
# building each case: a part it does not model at all (a 3-D, surface, stock or
# pie-of-pie chart) is a ValueError as its plot is built, a chart whose plot
# area holds no plot is an IndexError, and a bar chart grouped "standard" -- a
# value ECMA-376 allows and the library's table omits -- is a KeyError. Only
# the chart is given up; the deck around it still reads.
_UNREADABLE_CHART_ERRORS: Final = (IndexError, KeyError, ValueError)
# Each note stands alone and stays short, because the read tool's summary caps
# how long a single note may be and silently shortens one that runs past it.
_PPTX_NOTES: Final = (
    "A slide's words, tables and chart values are all in this text. What it "
    "cannot carry is the picture: where shapes sit, what arrows connect, what "
    "SmartArt lays out, and what the listed images show.",
    "So an answer that turns on how a slide is arranged, rather than on what "
    "it says, cannot be read out of this text. Say that instead of guessing "
    "the arrangement from the order the shapes are listed in.",
    "A chart's table is the data PowerPoint last saved with it, not a fresh "
    "read of a workbook it links to. Its categories are the labels the chart "
    "stores, so a date axis reads as Excel's serial number for each date.",
    "Animations, transitions, embedded video, audio and OLE objects, comments, "
    "and text that only the slide master or layout carries are not included.",
    "Shapes are read in the order PowerPoint stores them, back to front, which "
    "is not always the order they are read on screen. Shapes that are grouped "
    "are read where their group sits.",
    "Bullet characters are not resolved; a paragraph's outline level is shown "
    "as indentation instead. Speaker notes follow each slide's text under a "
    "Notes: line.",
    "A slide PowerPoint skips during a show is marked (hidden) in its heading "
    "and read like any other.",
)


def extract_pptx(source: BinaryIO, start_unit: int | None) -> ExtractedDocument:
    presentation = Presentation(source)
    slides = list(presentation.slides)
    total_units = len(slides)
    start = resolve_start_unit(start_unit, total_units=total_units, unit_kind="slide")
    budget = DocumentTextBudget(
        unit_kind="slide", total_units=total_units, start_unit=start
    )
    images: list[DocumentImage] = []
    charts: list[DocumentChart] = []
    outline: list[str] = []
    table_cut = False
    unread_charts = 0
    for number, slide in enumerate(slides[start - 1 :], start=start):
        location = f"slide {number}"
        title = slide.shapes.title
        heading = _slide_heading(slide, number, title)
        # The outline names slides; the markdown level belongs to the text.
        outline.append(heading.removeprefix("## "))
        lines: list[str] = []
        picture_parts: list[str] = []
        for shape in _flattened(slide.shapes):
            if shape == title:
                continue
            if isinstance(shape, Picture):
                picture_parts.extend(_picture_part(shape))
            elif isinstance(shape, GraphicFrame):
                if shape.has_chart:
                    entry = _chart_entry(shape.chart, location)
                    charts.append(entry)
                    tables, cut = _chart_tables(shape.chart)
                    table_cut = table_cut or cut
                    if not tables:
                        unread_charts += 1
                    lines.append(_chart_heading(shape.chart, entry.title))
                    lines.extend(tables)
                elif shape.has_table:
                    rows, cut = pipe_table(_table_rows(shape.table))
                    table_cut = table_cut or cut
                    lines.extend(rows)
            elif isinstance(shape, Shape):
                lines.extend(_shape_lines(shape))
        lines.extend(_notes_lines(slide))
        images.extend(
            DocumentImage(location=location, ref=part_name, count=count)
            for part_name, count in Counter(picture_parts).items()
        )
        # A slide the budget cannot hold ends this read; total_units still
        # counts every slide, and next_start_unit says where the rest begin.
        if not budget.add_unit(number, heading, *lines, ""):
            break
    return ExtractedDocument(
        document_format="pptx",
        unit_kind="slide",
        total_units=total_units,
        text=budget.finish(),
        content_truncated=budget.stopped or table_cut,
        next_start_unit=budget.next_start_unit,
        outline=tuple(outline),
        images=tuple(images),
        charts=tuple(charts),
        notes=(
            *_PPTX_NOTES,
            *_unread_chart_note(unread_charts),
            *table_cut_notes(table_cut),
            *budget.notes(),
        ),
    )


def _table_rows(table: Table) -> Iterator[list[str]]:
    """Each row's cell text, walking the table's rows once.

    ``Table.rows`` re-reads every row element on each index, so iterating it
    costs the square of the row count; the row elements are listed once here
    and handed to python-pptx's own row type.
    """

    rows = table.rows
    for tr in table._tbl.tr_lst:
        yield [_one_line(cell.text) for cell in _Row(tr, rows).cells]


def _unread_chart_note(unread_charts: int) -> tuple[str, ...]:
    """Say how many charts were listed without the numbers behind them."""

    if not unread_charts:
        return ()
    return (
        f"Data could not be read from {unread_charts} of these charts. An XY or "
        "bubble chart keeps its points outside the categories the library "
        "reports, and a 3-D, surface, stock or pie-of-pie chart it does not "
        "model at all.",
    )


def _flattened(shapes: Sequence[BaseShape]) -> Iterator[BaseShape]:
    """Every shape back to front, with a group replaced by the shapes inside it.

    Back to front is the order python-pptx returns and the order PowerPoint
    stores and selects shapes in: on a slide built from a layout it runs title
    first and then body, while ordering by position would depend on placeholder
    geometry inherited from that layout. Groups are expanded iteratively, so
    nesting cannot exhaust the interpreter's stack.
    """

    pending = list(reversed(shapes))
    while pending:
        shape = pending.pop()
        if isinstance(shape, GroupShape):
            pending.extend(reversed(list(shape.shapes)))
            continue
        yield shape


def _slide_heading(slide: Slide, number: int, title: Shape | None) -> str:
    name = _one_line(title.text_frame.text) if title is not None else ""
    # python-pptx exposes no property for whether PowerPoint shows a slide, so
    # the attribute is read from the element the library hands out.
    marker = " (hidden)" if slide.element.get("show") in _HIDDEN_SLIDE_VALUES else ""
    return f"## Slide {number}{marker}" + (f": {name}" if name else "")


def _shape_lines(shape: Shape) -> Iterator[str]:
    """A shape's paragraphs, indented by the outline level each one carries."""

    for paragraph in shape.text_frame.paragraphs:
        indent = "  " * min(paragraph.level, _MAX_OUTLINE_LEVEL)
        for line in paragraph.text.split(_LINE_BREAK):
            if line.strip():
                yield f"{indent}{line.strip()}"


def _notes_lines(slide: Slide) -> Iterator[str]:
    """The speaker notes, which carry what the slide itself does not say."""

    if not slide.has_notes_slide:
        return
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return
    lines = [
        line.strip()
        for line in frame.text.replace(_LINE_BREAK, "\n").split("\n")
        if line.strip()
    ]
    if lines:
        yield "Notes:"
        yield from lines


def _picture_part(picture: Picture) -> Iterator[str]:
    """The package part holding this picture, or nothing when it only links."""

    reference = picture.element.blip_rId
    if reference is not None:
        yield str(picture.part.related_part(reference).partname).lstrip("/")


def _chart_entry(chart: Chart, location: str) -> DocumentChart:
    """A chart's own title, empty when it shows a generated one or none."""

    title = chart.chart_title if chart.has_title else None
    text = title.text_frame.text if title is not None and title.has_text_frame else ""
    return DocumentChart(location=location, title=_one_line(text))


def _chart_heading(chart: Chart, title: str) -> str:
    """What the chart is, said above its data the way a slide says its own."""

    try:
        label = f"Chart ({chart.chart_type.name})"
    except _UNREADABLE_CHART_ERRORS:
        label = "Chart"
    return f"{label}: {title}" if title else label


def _chart_tables(chart: Chart) -> tuple[list[str], bool]:
    """A chart's saved data, one table per plot, and whether one was cut.

    A chart is the only place its numbers appear on a slide, so they are
    written out rather than left to the picture. Nothing comes back for a
    chart whose points python-pptx does not expose -- an XY or bubble plot
    reports no categories at all -- or one it cannot open; the caller counts
    those so the notes can say how many were listed without their data.

    The chart part is walked here rather than through python-pptx's plot,
    series and category sequences, because those size themselves by the point
    counts the file declares and re-read their elements on every index: a few
    kilobytes of chart XML could keep that busy for hours. Each element is read
    once, and only as many points and series as a table keeps.
    """

    lines: list[str] = []
    cut = False
    try:
        for x_chart in chart._chartSpace.chart.plotArea.xCharts:
            # Building the plot is what refuses a chart type python-pptx does
            # not model, which the caller reports as a chart without data.
            PlotFactory(x_chart, chart)
            labels = _category_labels(x_chart)
            series = [
                (_series_name(ser), _series_values(ser))
                # One past what a table shows, so a wider plot reads as cut.
                for ser in x_chart.sers[:MAX_TABLE_COLUMNS]
            ]
            if not labels or not series:
                continue
            rows, plot_cut = pipe_table(_plot_rows(labels, series))
            cut = cut or plot_cut
            if lines:
                # A combination chart is several plots in one part, and two
                # tables written back to back read as one table with a stray
                # header in the middle of it.
                lines.append("")
            lines.extend(rows)
    except _UNREADABLE_CHART_ERRORS:
        # python-pptx raises rather than passing over a plot it cannot build,
        # so the plot that raised ends this chart and the ones read before it
        # stand.
        return lines, cut
    return lines, cut


def _category_labels(x_chart: BaseChartElement) -> list[tuple[str, ...]]:
    """Each category's labels from the outermost level in, as python-pptx joins them.

    Labels come from the first series, which is where python-pptx reads them.
    A flat axis has a slot for every category the cache counts, empty where no
    label was saved; a hierarchical one has a row per saved leaf label, under
    the last parent label whose index does not pass the leaf's. At most
    ``MAX_TABLE_ROWS`` are read: with the header that is one row past what a
    table keeps, so a longer axis still reads as cut.
    """

    category_source = x_chart.cat
    if category_source is None:
        return []
    if category_source.multiLvlStrRef is None:
        count = min(_category_count(x_chart), MAX_TABLE_ROWS)
        points = x_chart.xpath("./c:ser[1]/c:cat//c:lvl[1]/c:pt") or x_chart.xpath(
            "./c:ser[1]/c:cat//c:pt"
        )
        # A repeated index keeps the last label saved for it, as python-pptx does.
        by_index = {point.idx: point for point in points if point.idx < count}
        return [(_point_label(by_index.get(index)),) for index in range(count)]
    levels = [level.pt_lst for level in category_source.lvls]
    if not levels:
        return []
    parent_levels: list[tuple[list[BaseOxmlElement], list[int]]] = []
    for level in levels[1:]:
        # An empty level ends the hierarchy, as it does in python-pptx.
        if not level:
            break
        # The highest index seen so far, so a bisection finds the first parent
        # past a leaf even when the file saves them out of order.
        parent_levels.append((level, list(accumulate((p.idx for p in level), max))))
    labels: list[tuple[str, ...]] = []
    for leaf in levels[0][:MAX_TABLE_ROWS]:
        path = [_point_label(leaf)]
        for level, highest_index in parent_levels:
            position = bisect_right(highest_index, leaf.idx)
            path.append(_point_label(level[max(position - 1, 0)]))
        labels.append(tuple(reversed(path)))
    return labels


def _category_count(x_chart: BaseChartElement) -> int:
    """How many categories the plot's axis counts, as the first series to say so.

    python-pptx asks one XPath for this across every series at once, which
    measured 16 s for 60,000 series; asking each series in document order
    returns the same count and usually stops at the first.
    """

    for ser in x_chart.iterchildren(qn("c:ser")):
        counts = ser.xpath(".//c:cat//c:ptCount")
        if counts:
            return int(counts[0].val)
    return 0


def _series_name(ser: CT_SeriesComposite) -> str:
    """The series' saved name, read where python-pptx reads it."""

    names = ser.xpath("./c:tx//c:pt/c:v/text()")
    return _one_line(str(names[0])) if names else ""


def _series_values(ser: CT_SeriesComposite) -> list[str]:
    """The first ``MAX_TABLE_ROWS`` values a series' cache counts, read in one pass.

    A slot without a saved point is empty. Where a cache saves one index twice
    the first point counts, as python-pptx reads it; a value that is not a
    number raises, which ends the chart's data.
    """

    caches = ser.xpath("./c:val")
    if not caches:
        return []
    cache = caches[0]
    count = min(cache.ptCount_val, MAX_TABLE_ROWS)
    by_index: dict[int, float] = {}
    for point in cache.xpath(".//c:pt"):
        if point.idx < count and point.idx not in by_index:
            by_index[point.idx] = point.value
    return [stored_value_text(by_index.get(index)) for index in range(count)]


def _point_label(point: BaseOxmlElement | None) -> str:
    # python-pptx turns a label element with no text into the string "None";
    # an empty label is what the chart shows there.
    return "" if point is None else (point.findtext(qn("c:v")) or "")


def _plot_rows(
    labels: Sequence[Sequence[str]], series: Sequence[tuple[str, Sequence[str]]]
) -> Iterator[Sequence[str]]:
    """Series across the header, one row per category, as the chart's sheet holds it.

    Categories outnumber series in every chart that is readable at a glance,
    and the table caps allow far more rows than columns, so the long axis runs
    down the rows. A hierarchical axis joins its levels into the one category
    cell. The category labels and each series' values are separate caches in
    the file, so a series shorter or longer than the categories pads rather
    than truncates.
    """

    yield ("Category", *(name for name, _ in series))
    yield from zip_longest(
        (" / ".join(label) for label in labels),
        *(values for _, values in series),
        fillvalue="",
    )


def _one_line(text: str) -> str:
    return " ".join(text.split())


__all__ = ["extract_pptx"]
