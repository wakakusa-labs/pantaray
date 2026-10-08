from __future__ import annotations

import base64
import datetime
import json
import re
import struct
import tracemalloc
import zipfile
import zlib
from io import BytesIO
from pathlib import Path
from typing import cast

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as docx_qn
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SheetImage
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches as SlideInches
from pypdf import PageObject, PdfWriter
from pypdf.generic import (
    ArrayObject,
    ContentStream,
    DictionaryObject,
    NameObject,
    NumberObject,
    StreamObject,
)

from pantaray_agents.local_runtime.tooling import documents as documents_package
from pantaray_agents.local_runtime.tooling.action_file_read_memory import (
    build_action_file_read_memory_input,
)
from pantaray_agents.local_runtime.tooling.documents import (
    MAX_DOCUMENT_BYTES,
    MAX_DOCUMENT_EXPANDED_BYTES,
    MAX_DOCUMENT_TEXT_CHARS,
    MAX_NOTEBOOK_OUTPUT_CHARS,
    MAX_SHEET_ROWS,
    extract_document,
)
from pantaray_agents.local_runtime.tooling.documents.document_model import (
    MAX_TABLE_COLUMNS,
    MAX_TABLE_ROWS,
    DocumentImage,
)
from pantaray_agents.local_runtime.tooling.tool_result_storage import (
    ACTION_TOOL_RESULT_INLINE_CHARACTER_LIMIT,
)
from pantaray_agents.schema.tool_result import serialize_json_tool_output
from pantaray_agents.tools.contract import BrokerPolicyError
from pantaray_agents.tools.files.read_output import (
    ReadToolOutput,
)

from .read_tool_broker_support import bootstrap_read_runtime_db, execute_read_tool

# A 1x1 PNG, so a fixture can embed a picture without a binary file in the repo.
PIXEL_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08"
    b"\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00"
    b"\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)
# Word writes a password-protected .docx as an OLE2 compound file, not a zip.
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def write_sample_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Quarterly review", level=1)
    document.add_paragraph("Revenue held steady across both regions.")
    document.add_heading("Risks", level=2)
    document.add_paragraph("Supply lead times", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "East"
    table.cell(1, 1).text = "120"
    document.add_picture(BytesIO(PIXEL_PNG), width=Inches(1))
    document.save(path)


@pytest.mark.asyncio
async def test_read_docx_returns_text_and_judgement_material(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_docx(context.workspace_path / "review.docx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "review.docx"}
    )

    output = outcome.output
    assert outcome.status == "success"
    assert output["kind"] == "document"
    assert output["document_format"] == "docx"
    assert output["unit_kind"] == "paragraph"
    assert output["total_units"] == 5
    assert output["outline"] == ["# Quarterly review", "## Risks"]
    assert "# Quarterly review" in output["content"]
    assert "- Supply lead times" in output["content"]
    assert "| Region | Revenue |" in output["content"]
    assert output["charts"] == []
    assert output["pages_without_text"] == []
    # A note the summary had to shorten ends in an ellipsis, which means the
    # note was written too long to survive the cap.
    assert output["notes"] and not any(note.endswith("…") for note in output["notes"])
    assert output["truncated"] is False
    assert output["truncation_reason"] is None
    assert output["next_offset"] is None
    images = output["images"]
    assert len(images) == 1
    assert images[0]["location"] == "paragraph 5"
    assert images[0]["ref"]
    assert images[0]["count"] == 1
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_read_docx_pages_on_the_text_file_cursor(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_docx(context.workspace_path / "review.docx")

    first = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "review.docx", "limit": 2},
    )
    assert first.output["kind"] == "document"
    assert first.output["offset"] == 1
    assert first.output["end_line"] == 2
    assert first.output["truncation_reason"] == "page_limit"
    assert first.output["next_offset"] == 3

    rest = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "review.docx", "offset": first.output["next_offset"]},
    )
    assert rest.output["next_offset"] is None
    assert rest.output["truncated"] is False
    assert first.output["content"] + rest.output["content"] == "".join(
        line + "\n"
        for line in (
            "# Quarterly review",
            "Revenue held steady across both regions.",
            "## Risks",
            "- Supply lead times",
            "| Region | Revenue |",
            "| --- | --- |",
            "| East | 120 |",
        )
    )
    # The summary describes the document, not the page, so it repeats unchanged.
    for field in ("total_units", "outline", "images", "notes"):
        assert rest.output[field] == first.output[field]


@pytest.mark.asyncio
async def test_extraction_budget_hands_back_the_paragraph_to_continue_from(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    document = Document()
    paragraph_count = MAX_DOCUMENT_TEXT_CHARS // 1_000 + 50
    for index in range(paragraph_count):
        document.add_paragraph(f"{index:04d}" + "x" * 996)
    document.save(context.workspace_path / "long.docx")

    first = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "long.docx"}
    )
    assert first.output["truncated"] is True
    assert first.output["total_units"] == paragraph_count
    assert any("character budget" in note for note in first.output["notes"])
    stopped_at = first.output["next_start_unit"]
    assert isinstance(stopped_at, int)

    last = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "long.docx", "offset": first.output["total_lines"]},
    )
    assert last.output["next_offset"] is None
    assert last.output["truncated"] is True
    assert last.output["truncation_reason"] == "document_budget"
    assert f"start_unit={stopped_at}" in str(last.output["retry_hint"])
    ReadToolOutput.model_validate(last.output)

    # The paragraph the budget stopped at is left out of the first read whole,
    # so the continuation starts with it: the two reads meet without a gap and
    # without repeating a paragraph.
    assert last.output["total_lines"] == stopped_at - 1
    assert str(last.output["content"]).startswith(f"{stopped_at - 2:04d}x")
    rest = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "long.docx", "start_unit": stopped_at},
    )
    assert str(rest.output["content"]).startswith(f"{stopped_at - 1:04d}x")
    assert rest.output["start_unit"] == stopped_at


@pytest.mark.asyncio
async def test_a_unit_larger_than_the_budget_is_cut_and_still_moves_on(
    tmp_path: Path,
) -> None:
    """The unit a read begins at is never dropped whole: that would never end."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    document = Document()
    document.add_paragraph("y" * (MAX_DOCUMENT_TEXT_CHARS + 1_000))
    document.add_paragraph("after the huge one")
    document.save(context.workspace_path / "huge.docx")

    first = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "huge.docx"}
    )
    assert first.output["truncated"] is True
    assert first.output["total_lines"] == 1
    assert any("longer on its own" in note for note in first.output["notes"])
    assert first.output["next_start_unit"] == 2

    rest = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "huge.docx", "start_unit": 2}
    )
    assert rest.output["content"] == "after the huge one\n"
    assert rest.output["next_start_unit"] is None
    assert rest.output["truncated"] is False


@pytest.mark.asyncio
async def test_image_summary_is_bounded_and_says_how_many_were_dropped(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    document = Document()
    for _ in range(60):
        document.add_picture(BytesIO(PIXEL_PNG), width=Inches(1))
    document.save(context.workspace_path / "gallery.docx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "gallery.docx"}
    )

    assert len(outcome.output["images"]) == 20
    assert any(
        "60 images" in note and "first 20" in note for note in outcome.output["notes"]
    )
    assert (
        len(serialize_json_tool_output(outcome.output))
        <= ACTION_TOOL_RESULT_INLINE_CHARACTER_LIMIT
    )


@pytest.mark.asyncio
async def test_damaged_docx_is_refused_rather_than_read_in_part(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "broken.docx").write_bytes(b"PK\x03\x04 truncated")

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "broken.docx"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNREADABLE"
    assert exc_info.value.fix_hint


@pytest.mark.asyncio
async def test_docx_that_expands_past_the_memory_limit_is_refused(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    source_path = context.workspace_path / "bomb.docx"
    write_sample_docx(source_path)
    chunk = b"\0" * (1024 * 1024)
    with zipfile.ZipFile(source_path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        with archive.open("word/media/padding.bin", "w") as entry:
            for _ in range(MAX_DOCUMENT_EXPANDED_BYTES // len(chunk) + 1):
                entry.write(chunk)
    # The archive stays small on disk, so only the declared expansion catches it.
    assert source_path.stat().st_size < MAX_DOCUMENT_BYTES

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "bomb.docx"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_TOO_LARGE"


@pytest.mark.asyncio
async def test_docx_cell_spanning_past_the_table_cap_is_cut_not_expanded(
    tmp_path: Path,
) -> None:
    """A cell's declared span is a count to stop at, not a row to build."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "wide"
    table.cell(0, 1).text = "after"
    span = OxmlElement("w:gridSpan")
    span.set(docx_qn("w:val"), str(2**31 - 1))
    table.cell(0, 0)._tc.get_or_add_tcPr().append(span)
    document.save(context.workspace_path / "span.docx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "span.docx"}
    )

    assert outcome.output["content"] == "".join(
        line + "\n"
        for line in (
            "| " + " | ".join(["wide"] * MAX_TABLE_COLUMNS) + " |",
            "| " + " | ".join(["---"] * MAX_TABLE_COLUMNS) + " |",
            f"[table cut to {MAX_TABLE_COLUMNS} columns and 1 rows]",
        )
    )
    assert outcome.output["truncated"] is True
    assert any("read cannot reach the rest" in note for note in outcome.output["notes"])
    assert "Part of this document was left out" in outcome.output["retry_hint"]


@pytest.mark.asyncio
async def test_docx_reads_text_inside_content_controls(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    document = Document()
    document.add_paragraph("before")
    block = document.add_paragraph("in a block control")._p
    sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    block.addprevious(sdt)
    sdt.append(content)
    content.append(block)
    inline_sdt, inline_content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    paragraph = document.add_paragraph("Name: ")
    paragraph._p.append(inline_sdt)
    inline_sdt.append(inline_content)
    inline_content.append(paragraph.add_run("Ada")._r)
    document.save(context.workspace_path / "form.docx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "form.docx"}
    )

    assert outcome.output["content"] == "before\nin a block control\nName: Ada\n"
    assert outcome.output["total_units"] == 3


@pytest.mark.asyncio
async def test_docx_vertical_merge_down_a_long_table_repeats_its_text(
    tmp_path: Path,
) -> None:
    """Every row of a merge reads the cell it continues, without walking back up."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    rows, columns = 500, 20
    document = Document()
    table = document.add_table(rows=rows, cols=columns)
    # Marked on the elements directly: python-docx's merge() is itself slow at
    # this size.
    for row_number, tr in enumerate(table._tbl.tr_lst):
        for tc in tr.tc_lst:
            tc.vMerge = "restart" if row_number == 0 else "continue"
    for column, cell in enumerate(table.rows[0].cells):
        cell.text = f"c{column}"
    document.save(context.workspace_path / "merged.docx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "merged.docx"}
    )

    merged_row = "| " + " | ".join(f"c{column}" for column in range(columns)) + " |"
    divider = "| " + " | ".join(["---"] * columns) + " |"
    # The table is longer than one read returns, so the window is checked: every
    # whole row after the divider is the merged row.
    lines = cast(str, outcome.output["content"]).split("\n")
    assert lines[:3] == [merged_row, divider, merged_row]
    assert set(lines[2:-1]) == {merged_row}


@pytest.mark.asyncio
async def test_protected_docx_is_refused_with_its_own_code(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "sealed.docx").write_bytes(OLE2_MAGIC + b"\x00" * 64)

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "sealed.docx"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_ENCRYPTED"
    assert exc_info.value.fix_hint


@pytest.mark.asyncio
async def test_legacy_word_document_is_refused_with_the_format_to_use(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "memo.doc").write_bytes(OLE2_MAGIC + b"\x00" * 64)

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "memo.doc"}
        )

    assert exc_info.value.code == "READ_LEGACY_DOCUMENT_UNSUPPORTED"
    assert ".docx" in (exc_info.value.fix_hint or "")


def write_sample_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Region", "Revenue", "Closed", "Active"])
    sheet.append(["East", 120, datetime.datetime(2026, 1, 2), True])
    sheet.append(["West", 0.5, datetime.datetime(2026, 2, 3, 9, 30), False])
    # Row 4 stays empty, and the merged range leaves B5 without a value.
    sheet["A5"] = "Total"
    sheet.merge_cells("A5:B5")
    sheet.add_image(SheetImage(BytesIO(PIXEL_PNG)), "D7")
    chart = BarChart()
    chart.title = "Revenue by region"
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True
    )
    sheet.add_chart(chart, "F2")
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "Draft"
    notes.sheet_state = "hidden"
    workbook.create_sheet("Empty")
    workbook.save(path)


def replace_package_part(path: Path, part: str, content: bytes) -> None:
    """Rewrite one part of an OOXML package, leaving the rest byte for byte."""

    rebuilt = BytesIO()
    with (
        zipfile.ZipFile(path) as source,
        zipfile.ZipFile(rebuilt, "w") as destination,
    ):
        for entry in source.infolist():
            data = content if entry.filename == part else source.read(entry.filename)
            destination.writestr(entry.filename, data)
    path.write_bytes(rebuilt.getvalue())


@pytest.mark.asyncio
async def test_read_xlsx_returns_a_table_per_sheet_and_judgement_material(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_xlsx(context.workspace_path / "book.xlsx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "book.xlsx"}
    )

    output = outcome.output
    assert output["kind"] == "document"
    assert output["document_format"] == "xlsx"
    assert output["unit_kind"] == "sheet"
    assert output["total_units"] == 3
    assert output["outline"] == ["Sales", "Notes (hidden)", "Empty"]
    assert output["content"] == "".join(
        line + "\n"
        for line in (
            "## Sheet: Sales",
            "| Region | Revenue | Closed | Active |",
            "| --- | --- | --- | --- |",
            "| East | 120 | 2026-01-02 | TRUE |",
            "| West | 0.5 | 2026-02-03 09:30:00 | FALSE |",
            "|  |  |  |  |",
            "| Total |  |  |  |",
            "",
            "## Sheet: Notes (hidden)",
            "| Draft |",
            "| --- |",
            "",
            "## Sheet: Empty",
        )
    )
    assert output["images"] == [{"location": "Sales", "ref": "D7", "count": 1}]
    assert output["charts"] == [{"location": "Sales!F2", "title": "Revenue by region"}]
    assert output["truncated"] is False
    assert output["truncation_reason"] is None
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_oversized_sheet_is_cut_without_stopping_the_next_sheet(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    workbook = Workbook()
    long_sheet = workbook.active
    long_sheet.title = "Long"
    for index in range(MAX_SHEET_ROWS + 10):
        long_sheet.append([index])
    workbook.create_sheet("After")["A1"] = "still read"
    workbook.save(context.workspace_path / "long.xlsx")

    first = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "long.xlsx"}
    )
    assert first.output["truncated"] is True
    assert any("were cut: Long" in note for note in first.output["notes"])

    # The cut bounds one sheet, so the sheets after it are still extracted.
    tail = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "long.xlsx", "offset": first.output["total_lines"] - 2},
    )
    assert tail.output["content"] == "## Sheet: After\n| still read |\n| --- |\n"
    assert tail.output["truncation_reason"] == "document_budget"


@pytest.mark.asyncio
async def test_xlsx_that_expands_entities_is_refused(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    source_path = context.workspace_path / "entities.xlsx"
    write_sample_xlsx(source_path)
    # Four nested entities turn 40 declared bytes into 10,000 expanded ones; the
    # same shape with more levels exhausts memory from a file that stays inside
    # every size cap, because the growth happens after decompression.
    replace_package_part(
        source_path,
        "xl/worksheets/sheet1.xml",
        b"""<?xml version="1.0"?>
<!DOCTYPE w [
<!ENTITY a "AAAAAAAAAA">
<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">
<!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">
<!ENTITY d "&c;&c;&c;&c;&c;&c;&c;&c;&c;&c;">
]>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&d;</t></is></c></row></sheetData>
</worksheet>""",
    )

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "entities.xlsx"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNREADABLE"


@pytest.mark.asyncio
async def test_xlsx_ranges_covering_the_whole_sheet_are_not_expanded(
    tmp_path: Path,
) -> None:
    """A merged range or a hyperlink names a range, not cells to build one by one."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    source_path = context.workspace_path / "ranges.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "kept"
    workbook.save(source_path)
    with zipfile.ZipFile(source_path) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    whole_sheet = (
        b'<mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells>'
        b'<hyperlinks><hyperlink ref="A1:XFD1048576" location="Sheet!A1"/>'
        b"</hyperlinks>"
    )
    replace_package_part(
        source_path,
        "xl/worksheets/sheet1.xml",
        sheet_xml.replace(b"</sheetData>", b"</sheetData>" + whole_sheet),
    )

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "ranges.xlsx"}
    )

    assert outcome.output["content"] == "## Sheet: Sheet\n| kept |\n| --- |\n"
    assert outcome.output["truncated"] is False


def test_xlsx_picture_drawn_many_times_is_not_read_once_per_anchor(
    tmp_path: Path,
) -> None:
    """Anchors are counted from the drawing; the image part is never copied per anchor."""

    # 1,000 x 1,000 pixels stored uncompressed: a 3 MB part that the package
    # deflates to a few kilobytes and counts once against the expansion cap.
    rows = b"".join(b"\x00" + b"\x00" * 3_000 for _ in range(1_000))
    header = struct.pack(">IIBBBBB", 1_000, 1_000, 8, 2, 0, 0, 0)
    picture = (
        b"\x89PNG\r\n\x1a\n"
        + png_chunk(b"IHDR", header)
        + png_chunk(b"IDAT", zlib.compress(rows, 0))
        + png_chunk(b"IEND", b"")
    )
    source_path = tmp_path / "pictures.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "kept"
    workbook.active.add_image(SheetImage(BytesIO(picture)), "C3")
    workbook.save(source_path)
    with zipfile.ZipFile(source_path) as archive:
        drawing_xml = archive.read("xl/drawings/drawing1.xml")
    anchor = re.search(rb"<oneCellAnchor>.*</oneCellAnchor>", drawing_xml, re.S)
    assert anchor is not None
    anchors = 200
    replace_package_part(
        source_path,
        "xl/drawings/drawing1.xml",
        drawing_xml.replace(anchor.group(0), anchor.group(0) * anchors),
    )

    tracemalloc.start()
    try:
        with source_path.open("rb") as source:
            document = extract_document(source=source, document_format="xlsx")
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert document.images == (
        DocumentImage(location="Sheet", ref="C3", count=anchors),
    )
    # One copy of the part per anchor would be 600 MB.
    assert peak < 32 * 1024 * 1024


def png_chunk(kind: bytes, data: bytes) -> bytes:
    return (
        struct.pack(">I", len(data))
        + kind
        + data
        + struct.pack(">I", zlib.crc32(kind + data))
    )


@pytest.mark.asyncio
async def test_damaged_xlsx_is_refused_rather_than_read_in_part(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "broken.xlsx").write_bytes(b"PK\x03\x04 truncated")

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "broken.xlsx"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNREADABLE"
    assert exc_info.value.fix_hint


@pytest.mark.asyncio
async def test_legacy_spreadsheet_is_refused_with_the_format_to_use(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "budget.xls").write_bytes(OLE2_MAGIC + b"\x00" * 64)

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "budget.xls"}
        )

    assert exc_info.value.code == "READ_LEGACY_DOCUMENT_UNSUPPORTED"
    assert ".xlsx" in (exc_info.value.fix_hint or "")


def write_sample_pptx(path: Path) -> None:
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly review"
    body = slide.placeholders[1].text_frame
    body.text = "Revenue held steady"
    bullet = body.add_paragraph()
    bullet.text = "Supply lead times"
    bullet.level = 1
    slide.shapes.add_picture(
        BytesIO(PIXEL_PNG), SlideInches(1), SlideInches(3), width=SlideInches(1)
    )
    table = slide.shapes.add_table(
        2, 2, SlideInches(1), SlideInches(4), SlideInches(4), SlideInches(1)
    ).table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "East"
    table.cell(1, 1).text = "120"
    slide.notes_slide.notes_text_frame.text = "Mention the hiring freeze"

    chart_data = CategoryChartData()
    chart_data.categories = ["East", "West"]
    chart_data.add_series("Revenue", (120.0, 80.0))
    chart = (
        presentation.slides.add_slide(blank_layout)
        .shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            SlideInches(1),
            SlideInches(1),
            SlideInches(4),
            SlideInches(3),
            chart_data,
        )
        .chart
    )
    chart.has_title = True
    chart.chart_title.text_frame.text = "Revenue by region"

    hidden = presentation.slides.add_slide(blank_layout)
    hidden.shapes.add_textbox(
        SlideInches(1), SlideInches(1), SlideInches(3), SlideInches(1)
    ).text_frame.text = "Draft"
    # PowerPoint keeps "skip this slide in a show" on the slide element itself.
    hidden.element.set("show", "0")

    presentation.save(path)


@pytest.mark.asyncio
async def test_read_pptx_returns_a_section_per_slide_and_judgement_material(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_pptx(context.workspace_path / "deck.pptx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "deck.pptx"}
    )

    output = outcome.output
    assert output["kind"] == "document"
    assert output["document_format"] == "pptx"
    assert output["unit_kind"] == "slide"
    assert output["total_units"] == 3
    assert output["outline"] == [
        "Slide 1: Quarterly review",
        "Slide 2",
        "Slide 3 (hidden)",
    ]
    assert output["content"] == "".join(
        line + "\n"
        for line in (
            "## Slide 1: Quarterly review",
            "Revenue held steady",
            "  Supply lead times",
            "| Region | Revenue |",
            "| --- | --- |",
            "| East | 120 |",
            "Notes:",
            "Mention the hiring freeze",
            "",
            "## Slide 2",
            "Chart (COLUMN_CLUSTERED): Revenue by region",
            "| Category | Revenue |",
            "| --- | --- |",
            "| East | 120.0 |",
            "| West | 80.0 |",
            "",
            "## Slide 3 (hidden)",
            "Draft",
        )
    )
    assert output["images"] == [
        {"location": "slide 1", "ref": "ppt/media/image1.png", "count": 1}
    ]
    assert output["charts"] == [{"location": "slide 2", "title": "Revenue by region"}]
    # A note the summary had to shorten ends in an ellipsis, which means the
    # note was written too long to survive the cap.
    assert output["notes"] and not any(note.endswith("…") for note in output["notes"])
    assert output["truncated"] is False
    assert output["truncation_reason"] is None
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_chart_categories_keep_their_hierarchy_and_missing_points(
    tmp_path: Path,
) -> None:
    """A category level and a point PowerPoint never saved both have to show."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    presentation = Presentation()
    chart_data = CategoryChartData()
    west = chart_data.categories.add_category("West")
    west.add_sub_category("SF")
    west.add_sub_category("LA")
    chart_data.categories.add_category("East").add_sub_category("NY")
    chart_data.add_series("Sales", (1.5, None, 3.0))
    presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        SlideInches(1),
        SlideInches(1),
        SlideInches(4),
        SlideInches(3),
        chart_data,
    )
    presentation.save(context.workspace_path / "cities.pptx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "cities.pptx"}
    )

    assert outcome.output["content"] == "".join(
        line + "\n"
        for line in (
            "## Slide 1",
            "Chart (COLUMN_CLUSTERED)",
            "| Category | Sales |",
            "| --- | --- |",
            "| West / SF | 1.5 |",
            "| West / LA |  |",
            "| East / NY | 3.0 |",
        )
    )


@pytest.mark.asyncio
async def test_a_chart_python_pptx_cannot_read_is_counted_not_fatal(
    tmp_path: Path,
) -> None:
    """A chart whose data is out of reach must not cost the deck around it."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    scatter_data = XyChartData()
    points = scatter_data.add_series("Trials")
    points.add_data_point(1.0, 2.0)
    points.add_data_point(2.0, 4.5)
    presentation.slides.add_slide(blank_layout).shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER,
        SlideInches(1),
        SlideInches(1),
        SlideInches(4),
        SlideInches(3),
        scatter_data,
    )
    column_data = CategoryChartData()
    column_data.categories = ["East", "West"]
    column_data.add_series("Revenue", (120.0, 80.0))
    solid = (
        presentation.slides.add_slide(blank_layout)
        .shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            SlideInches(1),
            SlideInches(1),
            SlideInches(4),
            SlideInches(3),
            column_data,
        )
        .chart
    )
    # PowerPoint writes a 3-D column chart as c:bar3DChart, a part python-pptx
    # models no plot for at all; it raises rather than skipping the chart.
    solid._chartSpace.xpath(".//c:barChart")[0].tag = qn("c:bar3DChart")
    presentation.save(context.workspace_path / "charts.pptx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "charts.pptx"}
    )

    output = outcome.output
    assert output["content"] == "## Slide 1\nChart (XY_SCATTER)\n\n## Slide 2\nChart\n"
    assert output["charts"] == [
        {"location": "slide 1", "title": ""},
        {"location": "slide 2", "title": ""},
    ]
    assert any(
        "Data could not be read from 2 of these charts" in note
        for note in cast(list[str], output["notes"])
    )
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_chart_point_counts_the_file_declares_bound_nothing_read(
    tmp_path: Path,
) -> None:
    """A cache that claims 2**31 - 1 points is read only as far as a table keeps."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    source_path = context.workspace_path / "counts.pptx"
    chart_data = CategoryChartData()
    chart_data.categories = ["East", "West"]
    chart_data.add_series("Revenue", (120.0, 80.0))
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        SlideInches(1),
        SlideInches(1),
        SlideInches(4),
        SlideInches(3),
        chart_data,
    )
    presentation.save(source_path)
    with zipfile.ZipFile(source_path) as archive:
        chart_xml = archive.read("ppt/charts/chart1.xml")
    replace_package_part(
        source_path,
        "ppt/charts/chart1.xml",
        re.sub(rb'<c:ptCount val="\d+"/>', b'<c:ptCount val="2147483647"/>', chart_xml),
    )

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "counts.pptx"}
    )

    empty_rows = MAX_TABLE_ROWS - 3
    assert outcome.output["content"] == "".join(
        line + "\n"
        for line in (
            "## Slide 1",
            "Chart (COLUMN_CLUSTERED)",
            "| Category | Revenue |",
            "| --- | --- |",
            "| East | 120.0 |",
            "| West | 80.0 |",
            *["|  |  |"] * empty_rows,
            f"[table cut to 2 columns and {MAX_TABLE_ROWS} rows]",
        )
    )


@pytest.mark.asyncio
async def test_long_hierarchical_chart_axis_is_cut_with_its_parents(
    tmp_path: Path,
) -> None:
    """Each leaf finds its parent label directly, however many the axis holds."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    leaves = 3_000
    chart_data = CategoryChartData()
    for index in range(leaves):
        chart_data.categories.add_category(f"P{index}").add_sub_category(f"L{index}")
    chart_data.add_series("Sales", [float(index) for index in range(leaves)])
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        SlideInches(1),
        SlideInches(1),
        SlideInches(4),
        SlideInches(3),
        chart_data,
    )
    presentation.save(context.workspace_path / "axis.pptx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "axis.pptx"}
    )

    kept = MAX_TABLE_ROWS - 1
    assert outcome.output["content"] == "".join(
        line + "\n"
        for line in (
            "## Slide 1",
            "Chart (COLUMN_CLUSTERED)",
            "| Category | Sales |",
            "| --- | --- |",
            *[f"| P{index} / L{index} | {float(index)} |" for index in range(kept)],
            f"[table cut to 2 columns and {MAX_TABLE_ROWS} rows]",
        )
    )


@pytest.mark.asyncio
async def test_paragraph_level_past_the_schema_indents_as_the_deepest_level(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(
        SlideInches(1), SlideInches(1), SlideInches(3), SlideInches(1)
    ).text_frame.paragraphs[0]
    paragraph.text = "deep"
    # python-pptx refuses to write a level past 8, so the attribute is set raw,
    # the way another producer could have saved it.
    paragraph._p.get_or_add_pPr().set("lvl", str(2**31 - 1))
    presentation.save(context.workspace_path / "deep.pptx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "deep.pptx"}
    )

    assert outcome.output["content"] == "## Slide 1\n" + "  " * 8 + "deep\n"


@pytest.mark.asyncio
async def test_text_inside_grouped_slide_shapes_is_read(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    inner = slide.shapes.add_textbox(
        SlideInches(1), SlideInches(1), SlideInches(3), SlideInches(1)
    )
    inner.text_frame.text = "inside a group"
    group = slide.shapes.add_group_shape([inner])
    nested = slide.shapes.add_textbox(
        SlideInches(1), SlideInches(3), SlideInches(3), SlideInches(1)
    )
    nested.text_frame.text = "inside a nested group"
    group.shapes.add_group_shape([nested])
    presentation.save(context.workspace_path / "grouped.pptx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "grouped.pptx"}
    )

    assert outcome.output["content"] == (
        "## Slide 1\ninside a group\ninside a nested group\n"
    )


@pytest.mark.asyncio
async def test_damaged_pptx_is_refused_rather_than_read_in_part(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "broken.pptx").write_bytes(b"PK\x03\x04 truncated")

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "broken.pptx"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNREADABLE"
    assert exc_info.value.fix_hint


@pytest.mark.asyncio
async def test_legacy_presentation_is_refused_with_the_format_to_use(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "pitch.ppt").write_bytes(OLE2_MAGIC + b"\x00" * 64)

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "pitch.ppt"}
        )

    assert exc_info.value.code == "READ_LEGACY_DOCUMENT_UNSUPPORTED"
    assert ".pptx" in (exc_info.value.fix_hint or "")


@pytest.mark.asyncio
async def test_document_read_becomes_read_memory(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_docx(context.workspace_path / "review.docx")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "review.docx"}
    )
    memory = build_action_file_read_memory_input(
        tool_id="read", status="completed", output=outcome.output
    )

    assert memory is not None
    assert memory.path == str(context.workspace_path / "review.docx")
    assert "Revenue held steady" in memory.content


def notebook_bytes(cells: list[dict[str, object]]) -> bytes:
    return json.dumps(
        {
            "cells": cells,
            "metadata": {"kernelspec": {"language": "python", "name": "python3"}},
            "nbformat": 4,
            "nbformat_minor": 5,
        }
    ).encode("utf-8")


def sample_notebook_cells() -> list[dict[str, object]]:
    return [
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": [
                "# Analysis\n",
                "Revenue held steady.\n",
                "\n",
                "```\n",
                "# not a heading\n",
                "```\n",
                "## Risks\n",
            ],
        },
        {
            "cell_type": "code",
            "execution_count": 1,
            "metadata": {},
            "source": ['print("hi")'],
            "outputs": [
                {"output_type": "stream", "name": "stdout", "text": ["hi\n"]},
                {
                    "output_type": "execute_result",
                    "execution_count": 1,
                    "data": {"text/plain": ["42"]},
                    "metadata": {},
                },
            ],
        },
        {
            "cell_type": "code",
            "execution_count": 2,
            "metadata": {},
            "source": ["boom()"],
            "outputs": [
                {
                    "output_type": "error",
                    "ename": "ValueError",
                    "evalue": "bad input",
                    # IPython colours a traceback, and the escapes are not content.
                    "traceback": ["[0;31mValueError[0m: bad input"],
                }
            ],
        },
        {
            "cell_type": "code",
            "execution_count": 3,
            "metadata": {},
            "source": ["plot()"],
            "outputs": [
                {
                    "output_type": "display_data",
                    "data": {
                        "image/png": base64.b64encode(PIXEL_PNG).decode("ascii"),
                        "text/html": "<div>" + "<span>x</span>" * 50 + "</div>",
                        "text/plain": ["<Figure size 640x480>"],
                    },
                    "metadata": {},
                }
            ],
        },
        {"cell_type": "raw", "metadata": {}, "source": ["raw content"]},
    ]


@pytest.mark.asyncio
async def test_read_ipynb_returns_a_section_per_cell_and_judgement_material(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "analysis.ipynb").write_bytes(
        notebook_bytes(sample_notebook_cells())
    )

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "analysis.ipynb"}
    )

    output = outcome.output
    assert output["kind"] == "document"
    assert output["document_format"] == "ipynb"
    assert output["unit_kind"] == "cell"
    assert output["total_units"] == 5
    assert output["outline"] == ["Cell 1: # Analysis", "Cell 1: ## Risks"]
    assert output["content"] == "".join(
        line + "\n"
        for line in (
            "## Cell 1 (markdown)",
            "# Analysis",
            "Revenue held steady.",
            "",
            "```",
            "# not a heading",
            "```",
            "## Risks",
            "",
            "## Cell 2 (code, In[1])",
            "```python",
            'print("hi")',
            "```",
            "stdout:",
            "hi",
            "Out[1]:",
            "42",
            "",
            "## Cell 3 (code, In[2])",
            "```python",
            "boom()",
            "```",
            "Error: ValueError: bad input",
            "ValueError: bad input",
            "",
            "## Cell 4 (code, In[3])",
            "```python",
            "plot()",
            "```",
            "Display:",
            "<Figure size 640x480>",
            "",
            "## Cell 5 (raw)",
            "raw content",
        )
    )
    assert output["images"] == [
        {"location": "cell 4", "ref": "cell 4 output 1 image/png", "count": 1}
    ]
    assert output["charts"] == []
    assert output["truncated"] is False
    assert output["truncation_reason"] is None
    # A note the summary had to shorten ends in an ellipsis, which means the
    # note was written too long to survive the cap.
    assert output["notes"] and not any(note.endswith("…") for note in output["notes"])
    assert any("last saved" in note for note in output["notes"])
    assert any("1 output representations" in note for note in output["notes"])
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_ipynb_result_carries_no_part_of_an_embedded_image(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "analysis.ipynb").write_bytes(
        notebook_bytes(sample_notebook_cells())
    )

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "analysis.ipynb"}
    )

    encoded = base64.b64encode(PIXEL_PNG).decode("ascii")
    assert encoded not in serialize_json_tool_output(outcome.output)


@pytest.mark.asyncio
async def test_one_long_notebook_output_is_cut_where_it_sits(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    lines = [f"step {index:04d} " + "x" * 40 + "\n" for index in range(120)]
    assert len("".join(lines)) > MAX_NOTEBOOK_OUTPUT_CHARS
    (context.workspace_path / "train.ipynb").write_bytes(
        notebook_bytes(
            [
                {
                    "cell_type": "code",
                    "execution_count": 1,
                    "metadata": {},
                    "source": ["train()"],
                    "outputs": [
                        {"output_type": "stream", "name": "stdout", "text": lines}
                    ],
                },
                {
                    "cell_type": "code",
                    "execution_count": 2,
                    "metadata": {},
                    "source": ["done()"],
                    "outputs": [],
                },
            ]
        )
    )

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "train.ipynb"}
    )

    output = outcome.output
    assert "step 0000" in output["content"]
    assert "step 0119" not in output["content"]
    assert (
        f"[output cut at {MAX_NOTEBOOK_OUTPUT_CHARS} characters]" in output["content"]
    )
    # The cut bounds one output, so the cell after it is still extracted.
    assert "## Cell 2 (code, In[2])" in output["content"]
    assert output["truncated"] is True
    assert output["truncation_reason"] == "document_budget"
    assert any("were longer than" in note for note in output["notes"])


@pytest.mark.asyncio
async def test_notebook_older_than_nbformat_4_is_refused_with_its_version(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    # nbformat 3 keeps cells inside worksheets, which this tool does not read.
    (context.workspace_path / "old.ipynb").write_bytes(
        json.dumps(
            {"nbformat": 3, "nbformat_minor": 0, "worksheets": [{"cells": []}]}
        ).encode("utf-8")
    )

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "old.ipynb"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNREADABLE"
    assert "nbformat 3" in str(exc_info.value)
    assert "save a fresh copy" in (exc_info.value.fix_hint or "")


@pytest.mark.asyncio
async def test_damaged_notebook_is_refused_rather_than_read_in_part(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "broken.ipynb").write_bytes(b'{"cells": [{"cell_ty')

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "broken.ipynb"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNREADABLE"
    assert exc_info.value.fix_hint


@pytest.mark.asyncio
async def test_deeply_nested_notebook_json_is_refused_without_crashing(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    # Past this depth json's scanner raises RecursionError, not a ValueError.
    depth = 100_000
    (context.workspace_path / "deep.ipynb").write_bytes(b"[" * depth + b"]" * depth)

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "deep.ipynb"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNREADABLE"


def write_sample_pdf(path: Path) -> None:
    """A PDF with a text page, a page with no text layer, and an image page.

    pypdf writes pages and bookmarks but draws nothing, so the text pages carry
    a content stream written here out of the same operators a PDF producer
    emits. Page 2 is left without one, which is what a scanned page looks like.
    """

    writer = PdfWriter()
    write_pdf_page(writer, ["Quarterly review", "Revenue held steady."])
    write_pdf_page(writer, [])
    third = write_pdf_page(writer, ["Risks"])
    third[NameObject("/Resources")][NameObject("/XObject")] = pdf_xobjects()
    chapter = writer.add_outline_item("Quarterly review", 0)
    writer.add_outline_item("Risks", 2, parent=chapter)
    writer.write(path)


def write_pdf_page(writer: PdfWriter, lines: list[str]) -> PageObject:
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = font
    resources = DictionaryObject()
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources
    if lines:
        drawn = " ".join(f"({line}) Tj 0 -16 Td" for line in lines)
        stream = ContentStream(None, None)
        stream.set_data(f"BT /F1 12 Tf 72 720 Td {drawn} ET".encode("ascii"))
        page.replace_contents(stream)
    return page


def pdf_xobjects() -> DictionaryObject:
    """An image beside a form, which is the pair the image count tells apart."""

    xobjects = DictionaryObject()
    for name, subtype in (("/Im0", "/Image"), ("/Fm0", "/Form")):
        entry = DictionaryObject()
        entry[NameObject("/Subtype")] = NameObject(subtype)
        xobjects[NameObject(name)] = entry
    return xobjects


def write_encrypted_pdf(path: Path, *, user_password: str) -> None:
    writer = PdfWriter()
    write_pdf_page(writer, ["Confidential"])
    writer.encrypt(
        user_password=user_password,
        owner_password="owner-secret",
        algorithm="AES-256",
    )
    writer.write(path)


@pytest.mark.asyncio
async def test_read_pdf_returns_text_and_judgement_material(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_pdf(context.workspace_path / "report.pdf")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "report.pdf"}
    )

    output = outcome.output
    assert outcome.status == "success"
    assert output["kind"] == "document"
    assert output["document_format"] == "pdf"
    assert output["unit_kind"] == "page"
    assert output["total_units"] == 3
    content = output["content"]
    assert isinstance(content, str)
    assert "## Page 1\nQuarterly review\nRevenue held steady." in content
    # A page with no text layer keeps its heading, so the numbering a renderer
    # is handed stays the document's own.
    assert "## Page 2\n\n## Page 3\nRisks" in content
    assert output["outline"] == ["Quarterly review (page 1)", "> Risks (page 3)"]
    assert output["pages_without_text"] == [2]
    assert output["images"] == [{"location": "page 3", "ref": "page 3", "count": 1}]
    assert output["charts"] == []
    assert output["truncated"] is False
    assert output["truncation_reason"] is None
    # A note the summary had to shorten ends in an ellipsis, which means the
    # note was written too long to survive the cap.
    assert output["notes"] and not any(note.endswith("…") for note in output["notes"])
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_read_pdf_opens_a_document_that_only_restricts_editing(
    tmp_path: Path,
) -> None:
    """An owner password blocks printing or editing, not reading."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_encrypted_pdf(context.workspace_path / "locked.pdf", user_password="")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "locked.pdf"}
    )

    assert outcome.output["kind"] == "document"
    assert "Confidential" in str(outcome.output["content"])


@pytest.mark.asyncio
async def test_password_protected_pdf_is_refused_with_its_own_code(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_encrypted_pdf(context.workspace_path / "sealed.pdf", user_password="letmein")

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "sealed.pdf"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_ENCRYPTED"


@pytest.mark.asyncio
async def test_text_saved_under_a_pdf_name_is_refused_rather_than_read(
    tmp_path: Path,
) -> None:
    """A .pdf path promises a PDF, the same way a .docx path promises a Word file."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "notes.pdf").write_text("plain notes\n", encoding="utf-8")

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "notes.pdf"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNREADABLE"


@pytest.mark.asyncio
async def test_pdf_without_an_extension_is_still_read_as_a_document(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_pdf(context.workspace_path / "report")

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "report"}
    )

    assert outcome.output["kind"] == "document"
    assert outcome.output["document_format"] == "pdf"


@pytest.mark.asyncio
async def test_pdf_extraction_gives_up_once_it_runs_past_its_time_limit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A PDF slower than the limit comes back partial, never as a whole document."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_pdf(context.workspace_path / "slow.pdf")
    monkeypatch.setattr(
        "pantaray_agents.local_runtime.tooling.documents.pdf."
        "MAX_PDF_EXTRACTION_SECONDS",
        -1.0,
    )

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "slow.pdf"}
    )

    output = outcome.output
    assert output["total_units"] == 3
    assert output["content"] == ""
    assert output["truncated"] is True
    assert output["truncation_reason"] == "document_budget"
    assert any("page 1 of 3" in note for note in cast(list[str], output["notes"]))
    # Nothing was read, so there is no page to continue from: pointing back at
    # page 1 would ask for the same slow read again, forever.
    assert output["next_start_unit"] is None


@pytest.mark.asyncio
async def test_pdf_page_redrawing_one_form_stops_at_the_time_limit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The limit holds inside a page, where pypdf parses a form again per draw."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    form = StreamObject()
    form.set_data(b"0 0 m\n" * 200_000)
    form.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Form"),
            NameObject("/BBox"): ArrayObject([NumberObject(0)] * 4),
            # pypdf skips a form with no resources rather than parsing it.
            NameObject("/Resources"): DictionaryObject(
                {NameObject("/ProcSet"): ArrayObject([NameObject("/PDF")])}
            ),
        }
    )
    form = form.flate_encode()
    resources = DictionaryObject()
    resources[NameObject("/XObject")] = DictionaryObject(
        {NameObject("/Fm0"): writer._add_object(form)}
    )
    page[NameObject("/Resources")] = resources
    drawing = ContentStream(None, None)
    drawing.set_data(b"/Fm0 Do\n" * 5_000)
    page.replace_contents(drawing)
    writer.write(context.workspace_path / "redraw.pdf")
    monkeypatch.setattr(
        "pantaray_agents.local_runtime.tooling.documents.pdf."
        "MAX_PDF_EXTRACTION_SECONDS",
        1.0,
    )

    outcome = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "redraw.pdf"}
    )

    output = outcome.output
    assert output["content"] == ""
    assert output["truncated"] is True
    assert any("page 1 of 1" in note for note in cast(list[str], output["notes"]))
    assert output["next_start_unit"] is None


@pytest.mark.asyncio
async def test_pdf_past_one_of_pypdfs_own_limits_is_refused_as_unreadable(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    # A CID width range that ends before it starts, which pypdf refuses with
    # LimitReachedError when it measures the font's glyphs.
    descendant = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/CIDFontType2"),
            NameObject("/BaseFont"): NameObject("/Broken"),
            NameObject("/W"): ArrayObject(
                [NumberObject(10), NumberObject(5), NumberObject(300)]
            ),
        }
    )
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type0"),
            NameObject("/BaseFont"): NameObject("/Broken"),
            NameObject("/Encoding"): NameObject("/Identity-H"),
            NameObject("/DescendantFonts"): ArrayObject(
                [writer._add_object(descendant)]
            ),
        }
    )
    resources = DictionaryObject()
    resources[NameObject("/Font")] = DictionaryObject({NameObject("/F1"): font})
    page[NameObject("/Resources")] = resources
    drawing = ContentStream(None, None)
    drawing.set_data(b"BT /F1 12 Tf 72 720 Td <0001> Tj ET")
    page.replace_contents(drawing)
    writer.write(context.workspace_path / "limit.pdf")

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "limit.pdf"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNREADABLE"


@pytest.mark.asyncio
async def test_extraction_out_of_memory_is_refused_as_too_large(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """An allocation a document sized past what memory holds refuses that read only."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_pdf(context.workspace_path / "report.pdf")

    def exhaust_memory(source: object, start_unit: object) -> None:
        raise MemoryError

    monkeypatch.setitem(documents_package._EXTRACTORS, "pdf", exhaust_memory)

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path, context=context, args={"path": "report.pdf"}
        )

    assert exc_info.value.code == "READ_DOCUMENT_TOO_LARGE"


@pytest.mark.asyncio
async def test_pdf_start_unit_skips_earlier_pages_without_extracting_them(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_pdf(context.workspace_path / "report.pdf")
    extracted: list[str] = []
    original = PageObject.extract_text

    def counted(page: PageObject, *args: object, **kwargs: object) -> str:
        text = str(original(page, *args, **kwargs))
        extracted.append(text.strip())
        return text

    monkeypatch.setattr(PageObject, "extract_text", counted)

    outcome = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "report.pdf", "start_unit": 3},
    )

    output = outcome.output
    assert str(output["content"]).startswith("## Page 3")
    assert "Quarterly review" not in str(output["content"])
    # Starting late has to be cheap, so the pages before it are never extracted.
    assert extracted == ["Risks"]
    assert output["start_unit"] == 3
    assert output["total_units"] == 3
    assert output["next_start_unit"] is None
    assert output["truncated"] is False
    # A bookmark is cheap and says where to read next, so the outline stays the
    # whole PDF's while images and pages_without_text cover what was extracted.
    assert output["outline"] == ["Quarterly review (page 1)", "> Risks (page 3)"]
    assert output["images"] == [{"location": "page 3", "ref": "page 3", "count": 1}]
    assert output["pages_without_text"] == []
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_pdf_read_in_budget_windows_covers_every_page_exactly_once(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A PDF past the text budget is read whole by following next_start_unit."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    writer = PdfWriter()
    for number in range(1, 7):
        write_pdf_page(writer, [f"Body {number}"])
    writer.write(context.workspace_path / "long.pdf")
    monkeypatch.setattr(
        "pantaray_agents.local_runtime.tooling.documents.document_model."
        "MAX_DOCUMENT_TEXT_CHARS",
        48,
    )

    read_lines: list[str] = []
    start_unit: object = 1
    reads = 0
    while start_unit is not None:
        reads += 1
        assert reads <= 6
        outcome = await execute_read_tool(
            db_path=db_path,
            context=context,
            args={"path": "long.pdf", "start_unit": start_unit},
        )
        assert outcome.output["total_units"] == 6
        read_lines.extend(
            line for line in str(outcome.output["content"]).splitlines() if line
        )
        start_unit = outcome.output["next_start_unit"]

    assert reads > 1
    assert read_lines == [
        line
        for number in range(1, 7)
        for line in (f"## Page {number}", f"Body {number}")
    ]


@pytest.mark.asyncio
async def test_start_unit_past_the_last_page_is_refused_with_the_count(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_pdf(context.workspace_path / "report.pdf")

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path,
            context=context,
            args={"path": "report.pdf", "start_unit": 4},
        )

    assert exc_info.value.code == "READ_DOCUMENT_UNIT_OUT_OF_RANGE"
    assert "3 pages" in str(exc_info.value)
    assert "start_unit" in (exc_info.value.fix_hint or "")


@pytest.mark.asyncio
async def test_start_unit_on_a_text_file_is_refused_rather_than_ignored(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "notes.md").write_text("one\ntwo\n", encoding="utf-8")

    with pytest.raises(BrokerPolicyError) as exc_info:
        await execute_read_tool(
            db_path=db_path,
            context=context,
            args={"path": "notes.md", "start_unit": 2},
        )

    assert exc_info.value.code == "READ_START_UNIT_UNSUPPORTED"
    assert "offset" in (exc_info.value.fix_hint or "")


@pytest.mark.asyncio
async def test_omitting_start_unit_reads_the_same_window_as_the_first_unit(
    tmp_path: Path,
) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_docx(context.workspace_path / "review.docx")

    default = await execute_read_tool(
        db_path=db_path, context=context, args={"path": "review.docx"}
    )
    first_unit = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "review.docx", "start_unit": 1},
    )

    assert default.output == first_unit.output
    assert default.output["start_unit"] == 1
    assert default.output["next_start_unit"] is None


@pytest.mark.asyncio
async def test_docx_start_unit_carries_the_tables_the_paragraph_owns(
    tmp_path: Path,
) -> None:
    """A table is read with the paragraph it follows, so neither is read twice."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_docx(context.workspace_path / "review.docx")

    outcome = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "review.docx", "start_unit": 3},
    )

    assert outcome.output["content"] == "".join(
        line + "\n"
        for line in (
            "## Risks",
            "- Supply lead times",
            "| Region | Revenue |",
            "| --- | --- |",
            "| East | 120 |",
        )
    )
    # The outline covers what was extracted, so the heading before the start is
    # not listed with the ones that are in this text.
    assert outcome.output["outline"] == ["## Risks"]

    after_the_table = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "review.docx", "start_unit": 5},
    )
    assert "| East | 120 |" not in str(after_the_table.output["content"])


@pytest.mark.asyncio
async def test_xlsx_start_unit_begins_at_that_sheet(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_xlsx(context.workspace_path / "book.xlsx")

    outcome = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "book.xlsx", "start_unit": 2},
    )

    output = outcome.output
    assert str(output["content"]).startswith("## Sheet: Notes (hidden)")
    assert "## Sheet: Sales" not in str(output["content"])
    assert output["outline"] == ["Notes (hidden)", "Empty"]
    assert output["total_units"] == 3
    assert output["images"] == []
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_pptx_start_unit_begins_at_that_slide(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_pptx(context.workspace_path / "deck.pptx")

    outcome = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "deck.pptx", "start_unit": 2},
    )

    output = outcome.output
    assert str(output["content"]).startswith("## Slide 2")
    assert "Quarterly review" not in str(output["content"])
    assert output["total_units"] == 3
    assert output["images"] == []
    assert output["charts"] == [{"location": "slide 2", "title": "Revenue by region"}]
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_ipynb_start_unit_begins_at_that_cell(tmp_path: Path) -> None:
    db_path, context = bootstrap_read_runtime_db(tmp_path)
    (context.workspace_path / "analysis.ipynb").write_bytes(
        notebook_bytes(sample_notebook_cells())
    )

    outcome = await execute_read_tool(
        db_path=db_path,
        context=context,
        args={"path": "analysis.ipynb", "start_unit": 2},
    )

    output = outcome.output
    assert str(output["content"]).startswith("## Cell 2 (code, In[1])")
    assert "# Analysis" not in str(output["content"])
    assert output["outline"] == []
    assert output["total_units"] == len(sample_notebook_cells())
    ReadToolOutput.model_validate(output)


@pytest.mark.asyncio
async def test_a_document_window_is_remembered_under_its_own_source(
    tmp_path: Path,
) -> None:
    """Two windows number their lines the same way, so the unit tells them apart."""

    db_path, context = bootstrap_read_runtime_db(tmp_path)
    write_sample_pdf(context.workspace_path / "report.pdf")

    windows = [
        build_action_file_read_memory_input(
            tool_id="read",
            status="completed",
            output=(
                await execute_read_tool(
                    db_path=db_path,
                    context=context,
                    args={"path": "report.pdf", "start_unit": start_unit},
                )
            ).output,
        )
        for start_unit in (1, 3)
    ]

    assert windows[0] is not None and windows[1] is not None
    assert windows[0].source_path.endswith("report.pdf#L1-L8")
    assert windows[1].source_path.endswith("report.pdf#page3#L1-L2")
