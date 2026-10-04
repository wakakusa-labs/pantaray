"""Real conversions through the real sandbox, with the LibreOffice on this Mac.

Opt-in (PANTARAY_LIBREOFFICE_IT=1) because it needs LibreOffice 26.2 or later
in /Applications, which CI does not have. It only runs conversions that succeed
under the verified profile: a conversion LibreOffice is refused something in
aborts with a crash dialog on the screen, so this is not a place to probe what
the sandbox denies.

The page counts are what the page renderer's addressing relies on: every slide
including hidden ones, and every sheet including hidden and empty ones.
"""

from __future__ import annotations

import os
import plistlib
import sys
from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pypdf import PdfReader

from pantaray_agents.local_runtime.tooling.documents.office_convert import (
    OfficeFormat,
    convert_office_to_pdf,
)

_LIBREOFFICE_APP = Path("/Applications/LibreOffice.app")
_MINIMUM_VERSION = (26, 2)


def _libreoffice_version() -> tuple[int, ...]:
    info_plist = _LIBREOFFICE_APP / "Contents" / "Info.plist"
    if not info_plist.is_file():
        return ()
    with info_plist.open("rb") as handle:
        version = str(plistlib.load(handle)["CFBundleShortVersionString"])
    return tuple(int(part) for part in version.split(".")[:2])


pytestmark = pytest.mark.skipif(
    os.environ.get("PANTARAY_LIBREOFFICE_IT") != "1"
    or sys.platform != "darwin"
    or _libreoffice_version() < _MINIMUM_VERSION,
    reason="set PANTARAY_LIBREOFFICE_IT=1 on macOS with LibreOffice 26.2+ installed",
)


def write_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("First page")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("Second page")
    document.save(str(path))


def write_pptx(path: Path) -> None:
    presentation = Presentation()
    layout = presentation.slide_layouts[5]
    for title in ("Shown", "Hidden", "Shown again"):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        if title == "Hidden":
            slide._element.set("show", "0")
    presentation.save(str(path))


def write_xlsx(path: Path) -> None:
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    for row in (("Quarter", "Sales"), ("Q1", 10), ("Q2", 14), ("Q3", 9)):
        data.append(row)
    hidden = workbook.create_sheet("Hidden")
    hidden["A1"] = "not printed"
    hidden.sheet_state = "hidden"
    workbook.create_sheet("Empty")
    chart = BarChart()
    chart.add_data(
        Reference(data, min_col=2, min_row=1, max_row=4), titles_from_data=True
    )
    chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=4))
    workbook.create_chartsheet("Chart").add_chart(chart)
    workbook.save(str(path))


@pytest.mark.parametrize(
    ("document_format", "write", "expected_pages"),
    [
        ("docx", write_docx, 2),
        # The hidden slide is exported too, so slide N stays page N.
        ("pptx", write_pptx, 3),
        # One page per sheet in workbook order, hidden, empty and chart sheets
        # included, so sheet N stays page N as well.
        ("xlsx", write_xlsx, 4),
    ],
)
async def test_each_format_converts_to_the_pages_its_units_address(
    tmp_path: Path,
    document_format: OfficeFormat,
    write: Callable[[Path], None],
    expected_pages: int,
) -> None:
    source = tmp_path / f"sample.{document_format}"
    write(source)
    destination = tmp_path / "sample.pdf"

    await convert_office_to_pdf(
        db_path=tmp_path / "runtime.db",
        libreoffice_app=_LIBREOFFICE_APP,
        source=source,
        document_format=document_format,
        destination=destination,
    )

    assert len(PdfReader(destination).pages) == expected_pages
